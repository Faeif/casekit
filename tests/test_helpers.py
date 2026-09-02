"""Shared test helpers, fixtures, and execution utilities for CaseKit tests."""

import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parent.parent
SKILLS = ROOT / "skills"
TEMPLATES = ROOT / "templates"
EXAMPLES = ROOT / "examples"
SCRIPTS = ROOT / "scripts"
FIXTURE_LAUNCH_EVENT = EXAMPLES / "launch-event"
CASEKIT_CLI = ROOT / "casekit.py"


def run_command(
    command: List[str],
    cwd: Optional[Path] = None,
    check: bool = True,
    input_text: Optional[str] = None,
) -> subprocess.CompletedProcess:
    """Execute a CLI command synchronously with captured stdout/stderr."""
    return subprocess.run(
        command,
        cwd=str(cwd) if cwd else str(ROOT),
        check=check,
        capture_output=True,
        text=True,
        input=input_text,
    )


def run_command_unchecked(
    command: List[str],
    cwd: Optional[Path] = None,
    input_text: Optional[str] = None,
) -> subprocess.CompletedProcess:
    """Execute a CLI command without raising on non-zero exit code."""
    return subprocess.run(
        command,
        cwd=str(cwd) if cwd else str(ROOT),
        check=False,
        capture_output=True,
        text=True,
        input=input_text,
    )


def create_temp_vault_copy(source_vault: Optional[Path] = None) -> Tuple[tempfile.TemporaryDirectory, Path]:
    """Create a temporary isolated copy of a fixture vault or empty workspace."""
    temp_dir = tempfile.TemporaryDirectory()
    temp_path = Path(temp_dir.name)
    vault_dest = temp_path / "test-vault"
    if source_vault and source_vault.exists():
        shutil.copytree(source_vault, vault_dest)
    else:
        vault_dest.mkdir(parents=True, exist_ok=True)
    return temp_dir, vault_dest


def build_mock_financial_model(
    output_path: Path,
    scenario_revenue: float = 1200000.0,
    gross_margin: float = 0.75,
    cash_runway_months: float = 18.0,
    cac_payback_months: float = 8.0,
    ltv_to_cac: float = 4.2,
) -> Path:
    """Create a test Excel model with standardized tabs and defined named ranges."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # Tab 1: 01_Assumptions
    ws_assump = wb.active
    ws_assump.title = "01_Assumptions"
    ws_assump["A1"] = "Assumption Driver"
    ws_assump["B1"] = "Value"
    ws_assump["A2"] = "Scenario"
    ws_assump["B2"] = "Base"
    ws_assump["A3"] = "SAFE_Investment_Amount"
    ws_assump["B3"] = 500000
    ws_assump["A4"] = "SAFE_Post_Money_Cap"
    ws_assump["B4"] = 10000000
    ws_assump["A5"] = "ESOP_Pool_Pct"
    ws_assump["B5"] = 0.15

    # Tab 2: 02_Unit_Economics
    ws_unit = wb.create_sheet(title="02_Unit_Economics")
    ws_unit["A1"] = "Metric"
    ws_unit["B1"] = "Value"
    ws_unit["A2"] = "Gross_Margin_Base"
    ws_unit["B2"] = gross_margin
    ws_unit["A3"] = "CAC_Payback_Months_Base"
    ws_unit["B3"] = cac_payback_months
    ws_unit["A4"] = "LTV_to_CAC_Base"
    ws_unit["B4"] = ltv_to_cac

    # Tab 3: 03_Three_Statements
    ws_stmt = wb.create_sheet(title="03_Three_Statements")
    ws_stmt["A1"] = "Line Item"
    ws_stmt["B1"] = "Base Scenario"
    ws_stmt["A2"] = "Gross_Revenue_Base"
    ws_stmt["B2"] = scenario_revenue
    ws_stmt["A3"] = "Cash_Runway_Months_Base"
    ws_stmt["B3"] = cash_runway_months
    ws_stmt["A4"] = "Ending_Cash_Base"
    ws_stmt["B4"] = scenario_revenue * 0.4

    # Tab 4: 04_Sensitivities
    ws_sens = wb.create_sheet(title="04_Sensitivities")
    ws_sens["A1"] = "Sensitivity Variable"
    ws_sens["B1"] = "Low"
    ws_sens["C1"] = "Base"
    ws_sens["D1"] = "High"
    ws_sens["A2"] = "Revenue"
    ws_sens["B2"] = scenario_revenue * 0.7
    ws_sens["C2"] = scenario_revenue
    ws_sens["D2"] = scenario_revenue * 1.4

    # Add Defined Names (Named Ranges)
    wb.defined_names.add(DefinedName("Gross_Revenue_Base", attr_text="'03_Three_Statements'!$B$2"))
    wb.defined_names.add(DefinedName("Gross_Margin_Base", attr_text="'02_Unit_Economics'!$B$2"))
    wb.defined_names.add(DefinedName("Cash_Runway_Months_Base", attr_text="'03_Three_Statements'!$B$3"))
    wb.defined_names.add(DefinedName("CAC_Payback_Months_Base", attr_text="'02_Unit_Economics'!$B$3"))
    wb.defined_names.add(DefinedName("LTV_to_CAC_Base", attr_text="'02_Unit_Economics'!$B$4"))
    wb.defined_names.add(DefinedName("Ending_Cash_Base", attr_text="'03_Three_Statements'!$B$4"))

    wb.save(output_path)
    return output_path


def send_mcp_jsonrpc_request(
    mcp_script_path: Path,
    request: Dict[str, Any],
) -> Dict[str, Any]:
    """Execute a single JSON-RPC request against an MCP server stdio process."""
    proc = subprocess.Popen(
        [sys.executable, str(mcp_script_path)],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    payload = json.dumps(request) + "\n"
    stdout, stderr = proc.communicate(input=payload, timeout=10)
    lines = [line.strip() for line in stdout.splitlines() if line.strip()]
    for line in lines:
        try:
            return json.loads(line)
        except json.JSONDecodeError:
            continue
    raise RuntimeError(f"No valid JSON-RPC response from MCP server. stdout: {stdout!r}, stderr: {stderr!r}")
